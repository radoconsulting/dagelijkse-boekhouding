"""Excel als tweede interface: exporteren, aanpassen in Excel en de wijzigingen terug inlezen.

Bewerkbaar in Excel (gele kolomkoppen):
- Transacties: Categorie (keuzelijst), Notitie en Tags. Een gesplitste transactie staat er per deel
  (id = txn_id#volgnummer); je past dan de categorie van dat deel aan.
- Budgetten: maandbudget en 'meenemen' per categorie
- Categorieën: nieuwe rijen (zonder id) en wijzigingen
- Regels: nieuwe rijen (zonder regel_id) en wijzigingen; uitschakelen via Actief = ONWAAR
- Rekeningen: naam, eigenaar, soort (zicht/spaar/kredietkaart) en een gekend saldo op een datum
Het tabblad Overzicht rekent met formules op de transacties, dus het volgt je aanpassingen meteen.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import duckdb
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .categoriseren import TEKENS, VELDEN, Regel, pas_regels_toe
from .beheer import _schone_tags, zet_notitie, zet_splitsdeel, zet_tags
from .db import BRON_EXCEL, categorie_id as maak_categorie_id, categorie_labels
from .utils import normaliseer_iban, parse_datum, schoon

SOORTEN = ("inkomst", "uitgave", "sparen", "transfer")
EURO = '#,##0.00 "€";[Red]-#,##0.00 "€"'
KOP = PatternFill("solid", fgColor="1F3A5F")
KOP_BEWERKBAAR = PatternFill("solid", fgColor="F2C14E")
GRIJS = Font(color="808080", size=9)


def _kop(ws, kolommen: list[tuple[str, int, bool]], rij: int = 1) -> None:
    for i, (naam, breedte, bewerkbaar) in enumerate(kolommen, start=1):
        c = ws.cell(row=rij, column=i, value=naam)
        c.fill = KOP_BEWERKBAAR if bewerkbaar else KOP
        c.font = Font(bold=True, color="1A1A1A" if bewerkbaar else "FFFFFF")
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = breedte
    ws.freeze_panes = ws.cell(row=rij + 1, column=1)


def _tabel(ws, naam: str, n_kolommen: int, n_rijen: int) -> None:
    ref = f"A1:{get_column_letter(n_kolommen)}{max(n_rijen, 1) + 1}"
    t = Table(displayName=naam, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)


def _keuzelijst(ws, bereik: str, formule: str, streng: bool = False) -> None:
    dv = DataValidation(type="list", formula1=formule, allow_blank=True)
    dv.errorStyle = "stop" if streng else "warning"
    dv.error = "Kies een waarde uit de lijst."
    ws.add_data_validation(dv)
    dv.add(bereik)


def exporteer_excel(con: duckdb.DuckDBPyConnection, pad: str | Path) -> Path:
    pad = Path(pad)
    pad.parent.mkdir(parents=True, exist_ok=True)
    labels = categorie_labels(con, alleen_actief=False)
    cats = con.execute(
        "SELECT categorie_id, hoofdcategorie, subcategorie, soort, actief FROM categorieen ORDER BY volgorde NULLS LAST, 1"
    ).fetchall()
    soort_per_label = {f"{h} / {s}": so for _, h, s, so, _ in cats}

    wb = Workbook()
    info = wb.active
    info.title = "Info"

    # --- Lijsten (verborgen) ---------------------------------------------------------
    lijsten = wb.create_sheet("Lijsten")
    lijsten.append(["Categorie", "Soort", "", "Veld", "Teken", "Soorten"])
    actieve_labels = [labels[c[0]] for c in cats if c[4]]
    n_lijst = max(len(actieve_labels), len(VELDEN), len(SOORTEN))
    for i in range(n_lijst):
        lab = actieve_labels[i] if i < len(actieve_labels) else None
        lijsten.append([
            lab, soort_per_label.get(lab) if lab else None, None,
            VELDEN[i] if i < len(VELDEN) else None,
            TEKENS[i] if i < len(TEKENS) else None,
            SOORTEN[i] if i < len(SOORTEN) else None,
        ])
    # Ook inactieve categorieën opzoekbaar voor de kolom Soort.
    for cid, h, s, so, actief in cats:
        if not actief:
            lijsten.append([f"{h} / {s}", so])
    lijsten.sheet_state = "hidden"
    cat_bereik = f"=Lijsten!$A$2:$A${len(actieve_labels) + 1}"

    # --- Transacties ----------------------------------------------------------------
    ws = wb.create_sheet("Transacties")
    kolommen = [
        ("id", 12, False), ("Datum", 11, False), ("Maand", 9, False), ("Rekening", 16, False),
        ("Bedrag", 12, False), ("Tegenpartij", 28, False), ("Omschrijving", 42, False), ("Mededeling", 30, False),
        ("Categorie", 40, True), ("Soort", 10, False), ("Bron", 10, False), ("Zekerheid", 9, False),
        ("Notitie", 30, True), ("Factuur", 12, False), ("Tags", 24, True),
    ]
    _kop(ws, kolommen)
    # Eén rij per boeking: een gesplitste transactie verschijnt per deel, zodat de formules juist optellen.
    rijen = con.execute(
        """SELECT b.boeking_id, b.boekingsdatum, strftime(b.boekingsdatum, '%Y-%m'), b.rekening_naam, b.bedrag,
                  coalesce(b.tegenpartij_naam, ''), coalesce(b.omschrijving, ''),
                  coalesce(b.mededeling, t.gestructureerde_mededeling, ''),
                  CASE WHEN b.categorie_id IS NULL THEN NULL ELSE b.categorie_label END,
                  b.categorie_bron, b.zekerheid, b.notitie, t.factuur_id, array_to_string(b.tags, ', ')
           FROM v_boekingen b JOIN transacties t USING (txn_id)
           ORDER BY b.boekingsdatum DESC, b.boeking_id"""
    ).fetchall()
    for r, rij in enumerate(rijen, start=2):
        # Zonder categorie telt een bedrag mee als inkomst of uitgave volgens het teken (zoals in de app).
        soort_formule = f'=IF(I{r}="",IF(E{r}>=0,"inkomst","uitgave"),IFERROR(VLOOKUP(I{r},Lijsten!$A:$B,2,FALSE),""))'
        waarden = [*rij[:9], soort_formule, *rij[9:]]  # kolom J (Soort) volgt de gekozen categorie
        for k, waarde in enumerate(waarden, start=1):
            ws.cell(row=r, column=k, value=waarde)
        ws.cell(row=r, column=1).font = GRIJS
        ws.cell(row=r, column=2).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=5).number_format = EURO
        ws.cell(row=r, column=12).number_format = "0%"
    n = len(rijen)
    _tabel(ws, "tblTransacties", len(kolommen), n)
    _keuzelijst(ws, f"I2:I{n + 500}", cat_bereik)

    # --- Overzicht (formules) -----------------------------------------------------------
    ov = wb.create_sheet("Overzicht", 1)
    maanden = [m for (m,) in con.execute(
        "SELECT DISTINCT strftime(boekingsdatum, '%Y-%m') FROM transacties ORDER BY 1"
    ).fetchall()][-24:]
    ov["A1"] = "Overzicht per maand"
    ov["A1"].font = Font(bold=True, size=14)
    ov["A2"] = "Formules op het tabblad Transacties: pas je daar een categorie aan, dan volgt dit overzicht. Negatief = geld dat buitengaat."
    ov["A2"].font = GRIJS
    ov.cell(row=4, column=1, value="Categorie").font = Font(bold=True)
    ov.cell(row=4, column=2, value="Soort").font = Font(bold=True)
    for j, m in enumerate(maanden, start=3):
        c = ov.cell(row=4, column=j, value=m)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="right")
        ov.column_dimensions[get_column_letter(j)].width = 12
    tot_kol = len(maanden) + 3
    ov.cell(row=4, column=tot_kol, value="Totaal").font = Font(bold=True)
    ov.column_dimensions["A"].width = 46
    ov.column_dimensions["B"].width = 10
    ov.column_dimensions[get_column_letter(tot_kol)].width = 13

    rij = 5
    totaalrijen = [
        ("Totaal inkomsten", "inkomst"), ("Totaal uitgaven", "uitgave"),
        ("Totaal sparen en beleggen", "sparen"), ("   waarvan niet gecategoriseerd", None),
    ]
    for titel, soort in totaalrijen:
        ov.cell(row=rij, column=1, value=titel).font = Font(bold=True)
        for j, _ in enumerate(maanden, start=3):
            col = get_column_letter(j)
            if soort:
                f = f'=SUMIFS(Transacties!$E:$E,Transacties!$J:$J,"{soort}",Transacties!$C:$C,{col}$4)'
            else:
                f = f'=SUMIFS(Transacties!$E:$E,Transacties!$I:$I,"=",Transacties!$C:$C,{col}$4)'
            ov.cell(row=rij, column=j, value=f).number_format = EURO
        ov.cell(row=rij, column=tot_kol, value=f"=SUM(C{rij}:{get_column_letter(tot_kol - 1)}{rij})").number_format = EURO
        rij += 1
    ov.cell(row=rij, column=1, value="Netto (inkomsten + uitgaven)").font = Font(bold=True)
    for j in range(3, tot_kol + 1):
        col = get_column_letter(j)
        ov.cell(row=rij, column=j, value=f"={col}5+{col}6").number_format = EURO
    rij += 2

    for cid, h, s, so, actief in cats:
        if so == "transfer":
            continue
        label = f"{h} / {s}"
        ov.cell(row=rij, column=1, value=label)
        ov.cell(row=rij, column=2, value=so).font = GRIJS
        for j, _ in enumerate(maanden, start=3):
            col = get_column_letter(j)
            ov.cell(row=rij, column=j, value=f"=SUMIFS(Transacties!$E:$E,Transacties!$I:$I,$A{rij},Transacties!$C:$C,{col}$4)").number_format = EURO
        ov.cell(row=rij, column=tot_kol, value=f"=SUM(C{rij}:{get_column_letter(tot_kol - 1)}{rij})").number_format = EURO
        rij += 1
    ov.freeze_panes = "C5"

    # --- Categorieën ------------------------------------------------------------------
    wc = wb.create_sheet("Categorieën")
    _kop(wc, [("categorie_id", 40, False), ("Hoofdcategorie", 24, True), ("Subcategorie", 36, True), ("Soort", 11, True), ("Actief", 8, True)])
    for cid, h, s, so, actief in cats:
        wc.append([cid, h, s, so, bool(actief)])
        wc.cell(row=wc.max_row, column=1).font = GRIJS
    _tabel(wc, "tblCategorieen", 5, len(cats))
    _keuzelijst(wc, f"D2:D{len(cats) + 200}", f"=Lijsten!$F$2:$F${len(SOORTEN) + 1}", streng=True)

    # --- Regels ----------------------------------------------------------------------
    wr = wb.create_sheet("Regels")
    _kop(wr, [("regel_id", 9, False), ("Prioriteit", 10, True), ("Naam", 26, True), ("Veld", 16, True), ("Patroon", 60, True),
              ("Teken", 8, True), ("Min bedrag", 11, True), ("Max bedrag", 11, True), ("Rekening", 20, True),
              ("Categorie", 40, True), ("Actief", 8, True), ("Bron", 10, False)])
    regels = con.execute(
        """SELECT regel_id, prioriteit, naam, veld, patroon, teken, min_bedrag, max_bedrag, rekening_iban, categorie_id, actief, bron
           FROM regels ORDER BY prioriteit, regel_id"""
    ).fetchall()
    for r in regels:
        wr.append([r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], labels.get(r[9], r[9]), bool(r[10]), r[11]])
        wr.cell(row=wr.max_row, column=1).font = GRIJS
    _tabel(wr, "tblRegels", 12, len(regels))
    extra = len(regels) + 300
    _keuzelijst(wr, f"D2:D{extra}", f"=Lijsten!$D$2:$D${len(VELDEN) + 1}", streng=True)
    _keuzelijst(wr, f"F2:F{extra}", f"=Lijsten!$E$2:$E${len(TEKENS) + 1}", streng=True)
    _keuzelijst(wr, f"J2:J{extra}", cat_bereik)

    # --- Rekeningen en facturen -------------------------------------------------------------
    wk = wb.create_sheet("Rekeningen")
    _kop(wk, [("IBAN", 22, False), ("Bank", 12, False), ("Naam", 28, True), ("Eigenaar", 20, True),
              ("Soort", 13, True), ("Gekend saldo", 14, True), ("Op datum", 12, True)])
    rek = con.execute(
        """SELECT iban, bank, naam, eigenaar, coalesce(soort, 'zicht'), startsaldo, startsaldo_datum
           FROM rekeningen ORDER BY bank, iban"""
    ).fetchall()
    for r in rek:
        wk.append(list(r))
        wk.cell(row=wk.max_row, column=6).number_format = EURO
        wk.cell(row=wk.max_row, column=7).number_format = "dd/mm/yyyy"
    _tabel(wk, "tblRekeningen", 7, len(rek))
    _keuzelijst(wk, f"E2:E{len(rek) + 50}", '"zicht,spaar,kredietkaart"', streng=True)

    wf = wb.create_sheet("Facturen")
    _kop(wf, [("Leverancier", 30, False), ("Factuurnr", 14, False), ("Datum", 11, False), ("Vervaldatum", 12, False),
              ("Bedrag", 12, False), ("OGM", 22, False), ("Status", 9, False), ("Omschrijving", 40, False), ("Bestand", 30, False)])
    fct = con.execute(
        """SELECT leverancier, factuurnummer, factuurdatum, vervaldatum, bedrag_totaal, gestructureerde_mededeling,
                  status, omschrijving, bestandsnaam FROM facturen ORDER BY factuurdatum DESC NULLS LAST"""
    ).fetchall()
    for r in fct:
        wf.append(list(r))
        wf.cell(row=wf.max_row, column=3).number_format = "dd/mm/yyyy"
        wf.cell(row=wf.max_row, column=4).number_format = "dd/mm/yyyy"
        wf.cell(row=wf.max_row, column=5).number_format = EURO
    _tabel(wf, "tblFacturen", 9, len(fct))

    # --- Budgetten -------------------------------------------------------------------------
    wbud = wb.create_sheet("Budgetten")
    _kop(wbud, [("categorie_id", 36, False), ("Categorie", 40, False), ("Maandbudget", 13, True), ("Meenemen", 10, True),
                ("Gemiddeld 6 mnd", 15, False), ("Voorstel", 10, False)])
    from .budget import budgetten as _budgetten, voorstel as _voorstel

    huidige_budgetten = {r.categorie_id: r for r in _budgetten(con).itertuples()}
    gemiddeld = {r.categorie_id: r for r in _voorstel(con).itertuples()}
    uitgave_cats = [c for c in cats if c[3] == "uitgave" and c[4]]
    for cid, h, s, _so, _actief in uitgave_cats:
        b, g = huidige_budgetten.get(cid), gemiddeld.get(cid)
        wbud.append([cid, f"{h} / {s}", b.maandbedrag if b else None, bool(b.meenemen) if b else False,
                     round(g.gemiddeld, 2) if g else None, g.voorstel if g else None])
        rij = wbud.max_row
        wbud.cell(row=rij, column=1).font = GRIJS
        for kol in (3, 5, 6):
            wbud.cell(row=rij, column=kol).number_format = EURO
    _tabel(wbud, "tblBudgetten", 6, len(uitgave_cats))

    # --- Info -------------------------------------------------------------------------
    info.column_dimensions["A"].width = 110
    teksten = [
        ("Dagelijkse boekhouding: Excel-export", Font(bold=True, size=14)),
        (f"Aangemaakt op {datetime.now():%d/%m/%Y %H:%M} met {n} transacties.", GRIJS),
        ("", None),
        ("Zo werk je ermee", Font(bold=True)),
        ("1. Kolommen met een gele kop mag je aanpassen. De andere kolommen worden bij het inlezen genegeerd.", None),
        ("2. Transacties: kies een categorie uit de keuzelijst, vul een notitie of tags in (komma-gescheiden, bv. vakantie, italië).", None),
        ("   Een lege categorie wist de categorie. Rijen met een id als abc#2 zijn delen van een gesplitste transactie.", None),
        ("   Budgetten: vul een maandbudget in (leeg of 0 = geen budget). Meenemen = over- of onderbesteding schuift door.", None),
        ("3. Categorieën: voeg een rij toe zonder categorie_id om een nieuwe categorie te maken.", None),
        ("4. Regels: voeg een rij toe zonder regel_id voor een nieuwe regel. Zet Actief op ONWAAR om een regel uit te schakelen.", None),
        ("   Patroon is een reguliere expressie (hoofdletter- en accentongevoelig), bv. colruyt|delhaize", None),
        ("   Veld: alles, tegenpartij_naam, tegenpartij_iban, omschrijving of mededeling. Teken: beide, in of uit.", None),
        ("5. Sla het bestand op en lees het in via de app (pagina Excel) of: python -m boekhouding excel-import <bestand>", None),
        ("", None),
        ("Wat je in Excel aanpast, telt als handmatige keuze en wordt nooit door regels of AI overschreven.", None),
        ("Laat de kolommen id, categorie_id en regel_id ongemoeid: zo weet de app welke rij welke is.", None),
    ]
    for i, (t, font) in enumerate(teksten, start=1):
        c = info.cell(row=i, column=1, value=t)
        if font:
            c.font = font

    wb.move_sheet("Lijsten", offset=len(wb.sheetnames))
    wb.active = 1
    wb.save(pad)
    return pad


@dataclass
class ExcelVerslag:
    transacties_gewijzigd: int = 0
    categorieen_nieuw: int = 0
    categorieen_gewijzigd: int = 0
    regels_nieuw: int = 0
    regels_gewijzigd: int = 0
    rekeningen_gewijzigd: int = 0
    budgetten_gewijzigd: int = 0
    opnieuw_gecategoriseerd: int = 0
    fouten: list[str] = field(default_factory=list)


def _rijen(ws) -> list[dict]:
    kop = [schoon(c.value) for c in ws[1]]
    uit = []
    for nr, rij in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v not in (None, "") for v in rij):
            continue
        d = {k: v for k, v in zip(kop, rij) if k}
        d["_rij"] = nr
        uit.append(d)
    return uit


def _bool(v) -> bool:
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() not in ("false", "onwaar", "0", "nee", "no", "faux", "")


def _getal(v):
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def importeer_excel(con: duckdb.DuckDBPyConnection, pad: str | Path) -> ExcelVerslag:
    wb = load_workbook(pad, data_only=False)
    v = ExcelVerslag()

    # 1. Categorieën eerst: nieuwe labels kunnen in transacties en regels gebruikt zijn.
    if "Categorieën" in wb.sheetnames:
        bestaand = {r[0]: r[1:] for r in con.execute("SELECT categorie_id, hoofdcategorie, subcategorie, soort, actief FROM categorieen").fetchall()}
        (max_volgorde,) = con.execute("SELECT coalesce(max(volgorde), 0) FROM categorieen").fetchall()[0]
        for r in _rijen(wb["Categorieën"]):
            h, s, so = schoon(r.get("Hoofdcategorie")), schoon(r.get("Subcategorie")), schoon(r.get("Soort"))
            actief = _bool(r.get("Actief", True))
            if not (h and s):
                continue
            so = (so or "uitgave").lower()
            if so not in SOORTEN:
                v.fouten.append(f"Categorieën rij {r['_rij']}: onbekende soort '{so}'.")
                continue
            cid = schoon(r.get("categorie_id"))
            if cid and cid in bestaand:
                if bestaand[cid] != (h, s, so, actief):
                    con.execute(
                        "UPDATE categorieen SET hoofdcategorie = ?, subcategorie = ?, soort = ?, actief = ? WHERE categorie_id = ?",
                        [h, s, so, actief, cid],
                    )
                    v.categorieen_gewijzigd += 1
            else:
                cid = maak_categorie_id(h, s)
                if cid in bestaand:
                    continue
                max_volgorde += 10
                con.execute("INSERT INTO categorieen VALUES (?, ?, ?, ?, ?, ?)", [cid, h, s, so, actief, max_volgorde])
                bestaand[cid] = (h, s, so, actief)
                v.categorieen_nieuw += 1

    label_naar_id = {lab: cid for cid, lab in categorie_labels(con, alleen_actief=False).items()}

    # 2. Rekeningen
    if "Rekeningen" in wb.sheetnames:
        for r in _rijen(wb["Rekeningen"]):
            iban = normaliseer_iban(r.get("IBAN"))
            if not iban:
                continue
            huidig = con.execute(
                "SELECT naam, eigenaar, coalesce(soort, 'zicht'), CAST(startsaldo AS DOUBLE), startsaldo_datum FROM rekeningen WHERE iban = ?",
                [iban],
            ).fetchall()
            if not huidig:
                continue
            soort = (schoon(r.get("Soort")) or huidig[0][2]).lower()
            if soort not in ("zicht", "spaar", "kredietkaart"):
                v.fouten.append(f"Rekeningen rij {r['_rij']}: onbekende soort '{soort}'.")
                continue
            saldo = _getal(r.get("Gekend saldo")) if "Gekend saldo" in r else huidig[0][3]
            datum = parse_datum(r.get("Op datum")) if "Op datum" in r else huidig[0][4]
            nieuw = (schoon(r.get("Naam")), schoon(r.get("Eigenaar")), soort, saldo, datum)
            if tuple(huidig[0]) != nieuw:
                con.execute(
                    "UPDATE rekeningen SET naam = ?, eigenaar = ?, soort = ?, startsaldo = ?, startsaldo_datum = ? WHERE iban = ?",
                    [*nieuw, iban],
                )
                v.rekeningen_gewijzigd += 1

    # 3. Regels
    regels_gewijzigd = False
    if "Regels" in wb.sheetnames:
        bestaand = {
            r[0]: r[1:]
            for r in con.execute(
                "SELECT regel_id, prioriteit, naam, veld, patroon, teken, min_bedrag, max_bedrag, rekening_iban, categorie_id, actief FROM regels"
            ).fetchall()
        }
        for r in _rijen(wb["Regels"]):
            patroon = schoon(r.get("Patroon"))
            label = schoon(r.get("Categorie"))
            if not patroon or not label:
                continue
            cid = label_naar_id.get(label)
            if cid is None:
                v.fouten.append(f"Regels rij {r['_rij']}: onbekende categorie '{label}'.")
                continue
            veld = (schoon(r.get("Veld")) or "alles").lower()
            teken = (schoon(r.get("Teken")) or "beide").lower()
            prioriteit = int(_getal(r.get("Prioriteit")) or 100)
            min_b, max_b = _getal(r.get("Min bedrag")), _getal(r.get("Max bedrag"))
            rekening = normaliseer_iban(r.get("Rekening"))
            actief = _bool(r.get("Actief", True))
            naam = schoon(r.get("Naam"))
            try:
                Regel(None, patroon, cid, veld=veld, teken=teken)
            except ValueError as e:
                v.fouten.append(f"Regels rij {r['_rij']}: {e}")
                continue
            rid = _getal(r.get("regel_id"))
            waarden = [prioriteit, naam, veld, patroon, teken, min_b, max_b, rekening, cid, actief]
            if rid is not None and int(rid) in bestaand:
                oud = bestaand[int(rid)]
                oud_vgl = [oud[0], oud[1], oud[2], oud[3], oud[4],
                           float(oud[5]) if oud[5] is not None else None,
                           float(oud[6]) if oud[6] is not None else None, oud[7], oud[8], oud[9]]
                if oud_vgl != waarden:
                    con.execute(
                        """UPDATE regels SET prioriteit = ?, naam = ?, veld = ?, patroon = ?, teken = ?, min_bedrag = ?,
                           max_bedrag = ?, rekening_iban = ?, categorie_id = ?, actief = ? WHERE regel_id = ?""",
                        [*waarden, int(rid)],
                    )
                    v.regels_gewijzigd += 1
                    regels_gewijzigd = True
            else:
                dubbel = any(
                    (o[2], o[3], o[4], o[8]) == (veld, patroon, teken, cid) for o in bestaand.values()
                )
                if dubbel:  # hetzelfde bestand twee keer ingelezen
                    continue
                con.execute(
                    """INSERT INTO regels (prioriteit, naam, veld, patroon, teken, min_bedrag, max_bedrag, rekening_iban,
                                           categorie_id, actief, bron) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'excel')""",
                    waarden,
                )
                bestaand[-(len(bestaand) + 1)] = tuple(waarden)
                v.regels_nieuw += 1
                regels_gewijzigd = True

    # 4. Transacties (en delen van gesplitste transacties): categorie, notitie en tags
    if "Transacties" in wb.sheetnames:
        huidig = {
            r[0]: r[1:]
            for r in con.execute(
                "SELECT boeking_id, txn_id, volgnr, categorie_id, notitie, array_to_string(tags, ', ') FROM v_boekingen"
            ).fetchall()
        }
        for r in _rijen(wb["Transacties"]):
            bid = schoon(r.get("id") or r.get("txn_id"))
            if not bid or bid not in huidig:
                continue
            txn_id, volgnr, oud_cat, oud_notitie, oud_tags = huidig[bid]
            label = schoon(r.get("Categorie"))
            if label and label not in label_naar_id:
                v.fouten.append(f"Transacties rij {r['_rij']}: onbekende categorie '{label}'.")
                continue
            nieuwe_cat = label_naar_id.get(label) if label else None
            notitie = schoon(r.get("Notitie"))
            tags = schoon(r.get("Tags")) if "Tags" in r else oud_tags
            gewijzigd = []
            if nieuwe_cat != oud_cat:
                gewijzigd.append("categorie_id")
            if notitie != (oud_notitie or None):
                gewijzigd.append("notitie")
            if _schone_tags(tags) != _schone_tags(oud_tags):
                gewijzigd.append("tags")
            if not gewijzigd:
                continue
            if volgnr is not None:
                zet_splitsdeel(con, txn_id, int(volgnr), categorie_id=nieuwe_cat, notitie=notitie, tags=tags,
                               velden=tuple(gewijzigd))
            else:
                if "categorie_id" in gewijzigd:
                    con.execute(
                        """UPDATE transacties SET categorie_id = ?, categorie_bron = ?, zekerheid = ?, regel_id = NULL,
                           ai_model = NULL, nagekeken = TRUE, gewijzigd_op = current_timestamp WHERE txn_id = ?""",
                        [nieuwe_cat, BRON_EXCEL if nieuwe_cat else None, 1.0 if nieuwe_cat else None, txn_id],
                    )
                if "notitie" in gewijzigd:
                    zet_notitie(con, txn_id, notitie)
                if "tags" in gewijzigd:
                    zet_tags(con, txn_id, tags)
            v.transacties_gewijzigd += 1

    # 5. Budgetten
    if "Budgetten" in wb.sheetnames:
        from .budget import bewaar_budget, budgetten as _budgetten

        huidige = {r.categorie_id: (r.maandbedrag, bool(r.meenemen)) for r in _budgetten(con).itertuples()}
        for r in _rijen(wb["Budgetten"]):
            cid = schoon(r.get("categorie_id"))
            if not cid or cid not in label_naar_id.values():
                continue
            bedrag = _getal(r.get("Maandbudget"))
            meenemen = _bool(r.get("Meenemen", False)) if r.get("Meenemen") is not None else False
            oud = huidige.get(cid)
            nieuw_ = (round(bedrag, 2), meenemen) if bedrag and bedrag > 0 else None
            if nieuw_ != (None if oud is None else (round(oud[0], 2), oud[1])):
                bewaar_budget(con, cid, bedrag, meenemen)
                v.budgetten_gewijzigd += 1

    if regels_gewijzigd:
        v.opnieuw_gecategoriseerd = pas_regels_toe(con)
    return v
