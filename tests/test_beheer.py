from datetime import date

import pytest
from conftest import lees
from hulp import rekening, txn
from openpyxl import load_workbook

from boekhouding import beheer, budget, rapporten
from boekhouding.categoriseren import pas_regels_toe
from boekhouding.db import nog_na_te_kijken
from boekhouding.excel import exporteer_excel, importeer_excel
from boekhouding.pipeline import importeer_bestand

SUPER = "boodschappen_supermarkt"
HUISHOUDEN = "boodschappen_drogisterij_en_huishouden"


@pytest.fixture
def ticket(con):
    rekening(con)
    return txn(con, date(2026, 10, 10), -100, naam="Colruyt", cat=SUPER, bron="regel")


def test_splitsen_en_rapporten(con, ticket):
    beheer.splits(con, ticket, [
        {"bedrag": -70, "categorie_id": SUPER},
        {"bedrag": -30, "categorie_id": HUISHOUDEN, "tags": "verbouwing, Verbouwing"},
    ])
    rijen = con.execute("SELECT boeking_id, bedrag, categorie_id, tags FROM v_boekingen ORDER BY boeking_id").fetchall()
    assert [(float(r[1]), r[2]) for r in rijen] == [(-70, SUPER), (-30, HUISHOUDEN)]
    assert rijen[1][0] == f"{ticket}#2" and rijen[1][3] == ["verbouwing"]
    per_sub = rapporten.per_categorie(con, niveau="sub").set_index("categorie")["bedrag"].to_dict()
    assert per_sub == {"Boodschappen / Supermarkt": 70, "Boodschappen / Drogisterij en huishouden": 30}
    # Regels raken een gesplitste transactie niet meer, en de oude categorie is weg.
    pas_regels_toe(con)
    assert con.execute("SELECT categorie_bron, categorie_id FROM transacties").fetchall() == [("gesplitst", None)]
    with pytest.raises(ValueError, match="gesplitst"):
        beheer.zet_categorie(con, ticket, SUPER)


@pytest.mark.parametrize("delen, fout", [
    ([{"bedrag": -70, "categorie_id": SUPER}, {"bedrag": -20, "categorie_id": HUISHOUDEN}], "verschil"),
    ([{"bedrag": -100, "categorie_id": SUPER}], "minstens twee"),
    ([{"bedrag": -130, "categorie_id": SUPER}, {"bedrag": 30, "categorie_id": HUISHOUDEN}], "teken"),
])
def test_ongeldige_splitsing(con, ticket, delen, fout):
    with pytest.raises(ValueError, match=fout):
        beheer.splits(con, ticket, delen)


def test_splitsing_ongedaan_maken(con, ticket):
    beheer.splits(con, ticket, [{"bedrag": -50, "categorie_id": SUPER}, {"bedrag": -50, "categorie_id": HUISHOUDEN}])
    beheer.verwijder_splitsing(con, ticket)
    # De standaardregel 'Supermarkten' zet de categorie terug.
    assert con.execute("SELECT categorie_id, categorie_bron FROM transacties").fetchall() == [(SUPER, "regel")]


def test_budget_telt_per_deel(con, ticket):
    budget.bewaar_budget(con, HUISHOUDEN, 50)
    beheer.splits(con, ticket, [{"bedrag": -60, "categorie_id": SUPER}, {"bedrag": -40, "categorie_id": HUISHOUDEN}])
    df = budget.budget_status(con, 2026, 10, date(2026, 10, 26)).set_index("categorie_id")
    assert df.loc[HUISHOUDEN, "besteed"] == 40


def test_tags(con, ticket):
    beheer.zet_tags(con, ticket, " Vakantie 2026 ,vakantie 2026, , Italië")
    assert con.execute("SELECT tags FROM transacties").fetchall()[0][0] == ["vakantie 2026", "italië"]
    assert beheer.alle_tags(con) == ["italië", "vakantie 2026"]
    t = rapporten.per_tag(con).set_index("tag")
    assert t.loc["italië", "uitgegeven"] == 100


def test_nakijken(con, inst):
    importeer_bestand(con, "kbc.csv", lees("kbc_voorbeeld.csv"), inst=inst)
    assert nog_na_te_kijken(con) == 11  # nieuwe imports zijn nog niet nagekeken
    tid = con.execute("SELECT txn_id FROM transacties WHERE bedrag = -25").fetchall()[0][0]
    beheer.zet_categorie(con, tid, "kinderen_jeugdbeweging_en_sport")
    assert nog_na_te_kijken(con) == 10
    alle = [r[0] for r in con.execute("SELECT txn_id FROM transacties").fetchall()]
    beheer.markeer_nagekeken(con, alle)
    assert nog_na_te_kijken(con) == 0


def test_excel_met_delen_tags_en_budgetten(con, ticket, tmp_path):
    beheer.splits(con, ticket, [{"bedrag": -70, "categorie_id": SUPER}, {"bedrag": -30, "categorie_id": SUPER}])
    pad = exporteer_excel(con, tmp_path / "x.xlsx")
    wb = load_workbook(pad)
    ws = wb["Transacties"]
    kop = [c.value for c in ws[1]]
    ids = [r[0].value for r in ws.iter_rows(min_row=2)]
    assert sorted(ids) == [f"{ticket}#1", f"{ticket}#2"]
    for r in ws.iter_rows(min_row=2):
        if r[0].value == f"{ticket}#2":
            r[kop.index("Categorie")].value = "Boodschappen / Drogisterij en huishouden"
            r[kop.index("Tags")].value = "verbouwing"
    wbud = wb["Budgetten"]
    bkop = [c.value for c in wbud[1]]
    for r in wbud.iter_rows(min_row=2):
        if r[0].value == SUPER:
            r[bkop.index("Maandbudget")].value = 450
            r[bkop.index("Meenemen")].value = True
    wb.save(tmp_path / "y.xlsx")

    v = importeer_excel(con, tmp_path / "y.xlsx")
    assert (v.transacties_gewijzigd, v.budgetten_gewijzigd, v.fouten) == (1, 1, [])
    assert con.execute("SELECT categorie_id, tags FROM splitsingen WHERE volgnr = 2").fetchall() == [(HUISHOUDEN, ["verbouwing"])]
    b = budget.budgetten(con).set_index("categorie_id")
    assert b.loc[SUPER, "maandbedrag"] == 450 and bool(b.loc[SUPER, "meenemen"])
    v2 = importeer_excel(con, tmp_path / "y.xlsx")
    assert (v2.transacties_gewijzigd, v2.budgetten_gewijzigd) == (0, 0)
