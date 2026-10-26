from conftest import lees
from openpyxl import load_workbook

from boekhouding.excel import exporteer_excel, importeer_excel
from boekhouding.pipeline import importeer_bestand


def _voorbereid(con, inst, tmp_path):
    for n in ("kbc_voorbeeld.csv", "belfius_voorbeeld.csv", "factuur_kinderopvang.pdf"):
        importeer_bestand(con, n, lees(n), inst=inst)
    return exporteer_excel(con, tmp_path / "export.xlsx")


def test_export_tabbladen(con, inst, tmp_path):
    wb = load_workbook(_voorbereid(con, inst, tmp_path))
    assert wb.sheetnames[:2] == ["Info", "Overzicht"]
    assert {"Transacties", "Categorieën", "Regels", "Rekeningen", "Facturen", "Lijsten"} <= set(wb.sheetnames)
    assert wb["Transacties"].max_row == 1 + 16
    assert wb["Lijsten"].sheet_state == "hidden"


def test_rondreis(con, inst, tmp_path):
    pad = _voorbereid(con, inst, tmp_path)
    wb = load_workbook(pad)
    ws = wb["Transacties"]
    kop = [c.value for c in ws[1]]
    for rij in ws.iter_rows(min_row=2):
        if rij[kop.index("Tegenpartij")].value == "VZW De Speelpleintjes":
            rij[kop.index("Categorie")].value = "Kinderen / Jeugdbeweging en sport"
            rij[kop.index("Notitie")].value = "Speelplein"
        if rij[kop.index("Bedrag")].value == -87.45:
            rij[kop.index("Categorie")].value = "Bestaat / Niet"
    wb["Categorieën"].append([None, "Vrije tijd", "Muziekles", "uitgave", True])
    wb["Regels"].append([None, 20, "Speelpleinen", "alles", "speelplein", "uit", None, None, None, "Kinderen / Jeugdbeweging en sport", True, None])
    wb.save(tmp_path / "bewerkt.xlsx")

    v = importeer_excel(con, tmp_path / "bewerkt.xlsx")
    assert (v.transacties_gewijzigd, v.categorieen_nieuw, v.regels_nieuw) == (1, 1, 1)
    assert len(v.fouten) == 1 and "Bestaat / Niet" in v.fouten[0]
    rij = con.execute("SELECT categorie_id, categorie_bron, notitie FROM transacties WHERE bedrag = -25").fetchall()[0]
    assert rij == ("kinderen_jeugdbeweging_en_sport", "excel", "Speelplein")
    assert con.execute("SELECT count(*) FROM categorieen WHERE subcategorie = 'Muziekles'").fetchall()[0][0] == 1

    # Hetzelfde bestand nogmaals inlezen verandert niets meer en maakt geen dubbele regels.
    v2 = importeer_excel(con, tmp_path / "bewerkt.xlsx")
    assert (v2.transacties_gewijzigd, v2.categorieen_nieuw, v2.regels_nieuw) == (0, 0, 0)
